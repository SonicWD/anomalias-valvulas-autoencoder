"""
Genera el Excel del plan de trabajo (estadísticas, filtro 3σ, lista de anómalas)
y los gráficos de dispersión requeridos.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage

BASE = Path(__file__).resolve().parent
DATA = BASE / "ACTIVIDAD_2"
OUT = BASE / "salidas"
OUT.mkdir(exist_ok=True)

SENSORS = [
    "Diametro_mm",
    "Peso_g",
    "Presion_bar",
    "Temperatura_C",
    "Dureza_HRC",
    "Tiempo_ciclo_s",
]


def main() -> None:
    train = pd.read_csv(DATA / "CasoEstudioValvulas_Entrenamiento.csv")
    lote = pd.read_csv(DATA / "CasoEstudioValvulas_LoteAInspeccionar.csv")

    # Resultados del autoencoder (si ya se corrió)
    ae_path = OUT / "resultados_lote_autoencoder.csv"
    ae = pd.read_csv(ae_path) if ae_path.exists() else None

    stats = train[SENSORS].describe().T
    stats = stats.rename(columns={"50%": "mediana"})
    stats["limite_inf_3std"] = stats["mean"] - 3 * stats["std"]
    stats["limite_sup_3std"] = stats["mean"] + 3 * stats["std"]

    # Filtro Excel sobre el lote
    lote_eval = lote.copy()
    fuera_cols = []
    for col in SENSORS:
        low = stats.loc[col, "limite_inf_3std"]
        high = stats.loc[col, "limite_sup_3std"]
        flag = (lote_eval[col] < low) | (lote_eval[col] > high)
        lote_eval[f"fuera_{col}"] = np.where(flag, "SI", "NO")
        fuera_cols.append(f"fuera_{col}")
    lote_eval["Anomalia_Excel"] = np.where(
        (lote_eval[fuera_cols] == "SI").any(axis=1), "SI", "NO"
    )
    lote_eval["n_sensores_fuera"] = (lote_eval[fuera_cols] == "SI").sum(axis=1)

    if ae is not None:
        merge_cols = ae[["ID_Valvula", "Error_reconstruccion", "Score_z", "Clasificacion"]]
        lote_eval = lote_eval.merge(merge_cols, on="ID_Valvula", how="left")
        lote_eval = lote_eval.rename(columns={"Clasificacion": "Clasificacion_AE"})

    # Gráficos de dispersión
    fig, ax = plt.subplots(figsize=(7, 5))
    ax.scatter(train["Temperatura_C"], train["Dureza_HRC"], s=14, alpha=0.5, c="#0F6E56")
    ax.set_xlabel("Temperatura (°C)")
    ax.set_ylabel("Dureza (HRC)")
    ax.set_title("Entrenamiento: Temperatura vs Dureza")
    fig.tight_layout()
    p1 = OUT / "excel_scatter_temp_dureza.png"
    fig.savefig(p1, dpi=140)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 5))
    ax.scatter(train["Diametro_mm"], train["Peso_g"], s=14, alpha=0.5, c="#378ADD")
    ax.set_xlabel("Diámetro (mm)")
    ax.set_ylabel("Peso (g)")
    ax.set_title("Entrenamiento: Diámetro vs Peso")
    fig.tight_layout()
    p2 = OUT / "excel_scatter_diametro_peso.png"
    fig.savefig(p2, dpi=140)
    plt.close(fig)

    # Workbook
    wb = Workbook()
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    red_fill = PatternFill("solid", fgColor="F4CCCC")
    green_fill = PatternFill("solid", fgColor="D9EAD3")

    def style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin

    # Hoja 1: Entrenamiento
    ws = wb.active
    ws.title = "Entrenamiento"
    for r in dataframe_to_rows(train, index=False, header=True):
        ws.append(r)
    style_header(ws)

    # Hoja 2: Estadísticas
    ws = wb.create_sheet("Estadisticas_Exploracion")
    stats_out = stats.reset_index().rename(columns={"index": "Sensor"})
    for r in dataframe_to_rows(stats_out, index=False, header=True):
        ws.append(r)
    style_header(ws)
    ws.append([])
    ws.append(["Nota"])
    ws.append(
        [
            "Los límites lim_inf/sup = media ± 3·desv.est. definen el rango normal univariado por sensor."
        ]
    )

    # Hoja 3: Lote con filtro Excel (+ AE si existe)
    ws = wb.create_sheet("Lote_Inspeccion")
    for r in dataframe_to_rows(lote_eval, index=False, header=True):
        ws.append(r)
    style_header(ws)
    # Colorear filas anómalas Excel
    col_idx = {cell.value: i + 1 for i, cell in enumerate(ws[1])}
    ae_col = col_idx.get("Clasificacion_AE")
    excel_col = col_idx.get("Anomalia_Excel")
    for row in range(2, ws.max_row + 1):
        if excel_col and ws.cell(row, excel_col).value == "SI":
            for c in range(1, ws.max_column + 1):
                ws.cell(row, c).fill = red_fill
        elif ae_col and ws.cell(row, ae_col).value == "ANOMALIA":
            for c in range(1, ws.max_column + 1):
                if ws.cell(row, c).fill.fgColor is None or ws.cell(row, c).fill.fgColor.rgb == "00000000":
                    ws.cell(row, c).fill = PatternFill("solid", fgColor="FCE5CD")

    # Hoja 4: Lista final anómalas
    ws = wb.create_sheet("Lista_Anomalas_Final")
    if ae is not None:
        anom = ae[ae["Clasificacion"] == "ANOMALIA"].copy()
        anom = anom.sort_values("Error_reconstruccion", ascending=False)
        # Unir info Excel
        anom = anom.merge(
            lote_eval[["ID_Valvula", "Anomalia_Excel", "n_sensores_fuera"]],
            on="ID_Valvula",
            how="left",
        )
        anom["Criterio"] = anom.apply(
            lambda r: "Ambos"
            if r["Anomalia_Excel"] == "SI"
            else "Solo Autoencoder (posible defecto correlacionado)",
            axis=1,
        )
        cols = [
            "ID_Valvula",
            "Error_reconstruccion",
            "Score_z",
            "Anomalia_Excel",
            "n_sensores_fuera",
            "Criterio",
        ] + SENSORS
        for r in dataframe_to_rows(anom[cols], index=False, header=True):
            ws.append(r)
        style_header(ws)
        for row in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(row, c).fill = red_fill
    else:
        ws.append(["Ejecutar primero 01_autoencoder_valvulas.py"])

    # Hoja 5: Respuesta conceptual Excel
    ws = wb.create_sheet("Analisis_Metodo_Excel")
    ws["A1"] = "¿El filtro univariado (media ± 3σ por sensor) detecta todas las anomalías reales?"
    ws["A1"].font = Font(bold=True)
    ws["A3"] = "Respuesta:"
    ws["A4"] = (
        "No. Ese método solo revisa cada sensor por separado. Puede escapársele un defecto "
        "multivariado: varios sensores se mueven juntos de forma coherente entre sí pero "
        "inconsistente con el proceso normal (por ejemplo, falla de tratamiento térmico que "
        "altera a la vez dureza, peso y presión sin que ninguno salga individualmente de su "
        "rango 3σ). El autoencoder sí captura esas correlaciones porque aprende el patrón "
        "conjunto de los seis sensores."
    )
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:F10")
    ws.column_dimensions["A"].width = 100

    # Hoja 6: Gráficos (imágenes)
    ws = wb.create_sheet("Graficos_Dispersion")
    ws["A1"] = "Gráficos de dispersión (exploración)"
    ws["A1"].font = Font(bold=True, size=12)
    try:
        img1 = XLImage(str(p1))
        img1.width, img1.height = 480, 340
        ws.add_image(img1, "A3")
        img2 = XLImage(str(p2))
        img2.width, img2.height = 480, 340
        ws.add_image(img2, "A22")
    except Exception as exc:
        ws["A3"] = f"No se pudieron insertar imágenes: {exc}"

    # Anchos
    for sheet in wb.worksheets:
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(22, max(12, len(str(col[0].value or "")) + 2))

    out_xlsx = OUT / "CasoEstudioValvulas_Entrega.xlsx"
    wb.save(out_xlsx)
    print(f"Excel generado: {out_xlsx}")


if __name__ == "__main__":
    main()
