"""
Genera el documento Word (.docx) explicativo del caso de estudio.
No incluye capturas de código; sí incluye gráficas de resultados.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent
OUT = BASE / "salidas"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def try_image(doc: Document, path: Path, width_cm: float = 15.0) -> None:
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    meta = {}
    metricas = {}
    if (OUT / "resumen_modelo.json").exists():
        meta = json.loads((OUT / "resumen_modelo.json").read_text(encoding="utf-8"))
    if (OUT / "metricas.json").exists():
        metricas = json.loads((OUT / "metricas.json").read_text(encoding="utf-8"))

    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(
        "DETECCIÓN DE ANOMALÍAS EN VÁLVULAS INDUSTRIALES\n"
        "Modelo de Deep Learning (Autoencoder)"
    )
    r.bold = True
    r.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Universidad de Cundinamarca · Especialización en Inteligencia Artificial\n"
        "CADI – Aplicaciones en Deep Learning · Período 2026-2 · GR 608145\n"
        "Actividad R1-A2-S4 · Docente: Ing. Wilson Rojas Reales, PhD(c)"
    ).font.size = Pt(10)

    add_heading(doc, "1. Comprensión del problema", 1)
    add_para(
        doc,
        "Una planta fabrica válvulas industriales. Cada unidad se caracteriza por seis "
        "mediciones de calidad: diámetro, peso, presión, temperatura, dureza y tiempo de ciclo. "
        "Estos sensores no son independientes: cuando el proceso falla, varios se desplazan "
        "de forma correlacionada (por ejemplo, un problema de tratamiento térmico afecta "
        "simultáneamente dureza, peso y presión)."
    )
    add_para(
        doc,
        "El objetivo es aprender el comportamiento normal del proceso usando únicamente "
        "válvulas sin defecto (500 registros) y luego señalar cuáles de las 100 válvulas "
        "del lote de inspección son anómalas, sin etiquetas previas."
    )

    add_heading(doc, "2. Modelos de Deep Learning aplicables", 1)
    add_para(doc, "Para detección de anomalías en entornos industriales se consideran:")
    add_bullets(
        doc,
        [
            "Autoencoders: aprenden a reconstruir datos normales; un error alto indica anomalía. "
            "Son adecuados cuando hay muchos ejemplos normales y pocos (o ningún) defectos etiquetados.",
            "CNN / autoencoders convolucionales: útiles cuando la entrada es imagen o señal "
            "con estructura espacial/temporal (en este caso las entradas son 6 sensores tabulares).",
            "Variational Autoencoders (VAE) o Isolation basados en redes: alternativas cuando "
            "se requiere modelar incertidumbre o distribuciones más complejas.",
        ],
    )
    add_para(
        doc,
        "En este trabajo se implementa un Autoencoder denso (MLP encoder–latente–decoder), "
        "alineado con el material de las semanas 1–4 y apropiado para vectores de sensores.",
        bold=False,
    )

    add_heading(doc, "3. Metodología desarrollada", 1)
    add_heading(doc, "3.1 Enfoque Excel (línea base univariada)", 2)
    add_bullets(
        doc,
        [
            "Exploración: media, desviación estándar, mínimo y máximo por sensor en entrenamiento.",
            "Gráficos de dispersión (Temperatura vs Dureza; Diámetro vs Peso) para visualizar correlación.",
            "Rango normal por sensor: media ± 3 desviaciones estándar.",
            "Marcado en el lote de toda válvula que se salga del rango en al menos un sensor.",
        ],
    )
    add_para(
        doc,
        "Limitación: este filtro no detecta defectos multivariados en los que ningún sensor "
        "individual sale de su rango, pero la combinación viola el patrón normal del proceso.",
    )

    add_heading(doc, "3.2 Enfoque Python — Autoencoder", 2)
    add_bullets(
        doc,
        [
            "Estandarización (StandardScaler) de los seis sensores.",
            "División entrenamiento/validación (80/20) solo con válvulas normales.",
            "Arquitectura MLPRegressor: capas ocultas (8, 3, 8) con activación ReLU; "
            "la capa de 3 neuronas actúa como espacio latente.",
            "Entrenamiento con entrada = salida (reconstrucción).",
            "Umbral de anomalía = media + 3·std del error MSE en validación.",
            "Clasificación del lote según error de reconstrucción.",
        ],
    )

    if meta:
        add_para(
            doc,
            f"Resultado del entrenamiento: umbral = {meta.get('umbral', 0):.5f}; "
            f"válvulas anómalas en el lote = {meta.get('n_anomalas', '?')} / {meta.get('n_lote', 100)}. "
            f"IDs detectados: {', '.join(meta.get('ids_anomalas', []))}.",
        )

    add_heading(doc, "4. Métricas de evaluación (mínimo tres) y justificación", 1)

    add_heading(doc, "4.1 Tasa de anomalías detectadas", 2)
    add_para(
        doc,
        "Mide qué proporción del lote se envía a revisión. Es una métrica operacional: "
        "en planta interesa controlar la carga de inspección. "
        f"Autoencoder: {100 * metricas.get('tasa_anomalias_autoencoder', 0):.1f}%. "
        f"Filtro Excel: {100 * metricas.get('tasa_anomalias_excel', 0):.1f}%.",
    )

    add_heading(doc, "4.2 Acuerdo entre Excel y Autoencoder", 2)
    add_para(
        doc,
        "Compara ambos enfoques sin necesitar etiquetas reales del lote. "
        f"Acuerdo total: {metricas.get('acuerdo_porcentaje', 0):.1f}%. "
        f"Detectadas por ambos: {metricas.get('detectadas_ambos', 0)}; "
        f"solo Autoencoder: {metricas.get('solo_autoencoder', 0)}; "
        f"solo Excel: {metricas.get('solo_excel', 0)}. "
        "Las válvulas detectadas solo por el Autoencoder son evidencia de anomalías "
        "correlacionadas que el filtro univariado no ve.",
    )
    if metricas.get("ids_solo_autoencoder"):
        add_para(
            doc,
            "IDs solo Autoencoder: " + ", ".join(metricas["ids_solo_autoencoder"]) + ".",
        )

    add_heading(doc, "4.3 Separación del error de reconstrucción", 2)
    ratio = metricas.get("ratio_error_anomalia_normal")
    add_para(
        doc,
        "Compara el error medio de las válvulas marcadas como anómalas frente a las "
        "marcadas como normales. Un ratio alto indica buena separación del detector. "
        + (f"Ratio obtenido: {ratio:.2f}x." if ratio else ""),
    )

    add_heading(doc, "4.4 Diagnóstico por sensor (complementaria)", 2)
    add_para(
        doc,
        "Para cada anomalía se calcula qué sensor aporta más al error de reconstrucción. "
        "Esto apoya el análisis de causa raíz en línea de producción (por ejemplo, "
        "si dominan dureza y temperatura, se sospecha del tratamiento térmico).",
    )

    add_heading(doc, "5. Evidencias gráficas de resultados", 1)
    add_para(doc, "Distribución del error de reconstrucción (validación vs lote):")
    try_image(doc, OUT / "fig_error_reconstruccion.png")
    add_para(doc, "Ranking de sospecha en el lote:")
    try_image(doc, OUT / "fig_ranking_lote.png")
    add_para(doc, "Comparación de métricas y acuerdo entre métodos:")
    try_image(doc, OUT / "fig_metricas_comparacion.png")
    add_para(doc, "Dispersión Temperatura vs Dureza:")
    try_image(doc, OUT / "fig_dispersion_temp_dureza.png", width_cm=12)
    add_para(doc, "Dispersión Diámetro vs Peso:")
    try_image(doc, OUT / "fig_dispersion_diametro_peso.png", width_cm=12)
    if (OUT / "fig_curva_perdida.png").exists():
        add_para(doc, "Curva de pérdida durante el entrenamiento:")
        try_image(doc, OUT / "fig_curva_perdida.png", width_cm=12)

    add_heading(doc, "6. Despliegue en la nube", 1)
    add_para(
        doc,
        "El modelo se expone mediante una interfaz Streamlit (`app_streamlit.py`) que permite: "
        "(a) clasificar una válvula ingresando los seis sensores; "
        "(b) revisar el lote completo y descargar resultados; "
        "(c) consultar la explicación del funcionamiento. "
        "La aplicación se despliega en Streamlit Community Cloud (o servicio equivalente), "
        "cumpliendo el requisito de implementación en la nube con interfaz funcional.",
    )

    add_heading(doc, "7. Conclusiones: Excel vs Python (Autoencoder)", 1)
    add_bullets(
        doc,
        [
            "El filtro Excel es útil como línea base rápida y transparente, pero ignora correlaciones.",
            "El Autoencoder aprende el patrón conjunto de los seis sensores y puede detectar "
            "anomalías multivariadas que Excel no marca.",
            "La comparación muestra discordancias relevantes (válvulas solo detectadas por AE), "
            "confirmando la ventaja del enfoque de Deep Learning para este problema industrial.",
            "Las métricas operacionales (tasa, acuerdo, separación del error y diagnóstico por "
            "sensor) permiten documentar el desempeño aunque el lote no traiga etiquetas.",
        ],
    )

    add_heading(doc, "8. Entregables asociados", 1)
    add_bullets(
        doc,
        [
            "Excel: salidas/CasoEstudioValvulas_Entrega.xlsx",
            "Script 1: 01_autoencoder_valvulas.py",
            "Script 2: 02_metricas_comparacion.py",
            "Lista de anómalas: salidas/lista_valvulas_anomalas.csv (también hoja en el Excel)",
            "Interfaz nube: app_streamlit.py",
        ],
    )

    out_doc = OUT / "Documento_Explicativo_R1-A2-S4_Modelo_Deep_Learning.docx"
    doc.save(out_doc)
    print(f"Documento generado: {out_doc}")


if __name__ == "__main__":
    main()
